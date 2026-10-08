from __future__ import annotations

import re
from pathlib import Path
from openpyxl import load_workbook

from .models import Workflow

ALIASES = {
    "id": {"id", "workflow id", "workflow_id"},
    "name": {"name", "workflow", "workflow name", "workflow_name", "process"},
    "description": {"description", "purpose", "summary", "trigger"},
    "inputs": {"input", "inputs", "required input", "required inputs"},
    "steps": {"steps", "required steps", "workflow steps", "process steps"},
    "tools": {"tools", "tools/apis", "tools/api", "tools/apis required", "tools_required", "apis", "api"},
    "conditions": {"conditions", "condition", "decisions", "conditions/decisions", "decision logic", "decision_logic", "rules"},
    "expected_output": {"expected output", "expected_output", "output", "result", "final output"},
}


def _norm(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _items(value: object) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    return [re.sub(r"^\s*(?:\d+[.)]|[-•])\s*", "", x).strip() for x in re.split(r";|\n|→|->", text) if x.strip()]


def _field(title: object) -> str | None:
    value = _norm(title)
    for canonical, aliases in ALIASES.items():
        if value in aliases:
            return canonical
    return None


def _workflow(values: dict[str, object], fallback_name: str) -> Workflow | None:
    name = str(values.get("name") or fallback_name).strip()
    steps = _items(values.get("steps"))
    # Ignore empty/template rows but allow a workflow with just a name and description.
    if not name or (not steps and not values.get("description")):
        return None
    slug = str(values.get("id") or re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-"))
    return Workflow(id=slug, name=name, description=str(values.get("description") or ""),
                    inputs=_items(values.get("inputs")), steps=steps, tools=_items(values.get("tools")),
                    conditions=_items(values.get("conditions")), expected_output=str(values.get("expected_output") or ""))


def load_workflows(path: str | Path) -> list[Workflow]:
    """Read row-based or field/value worksheet workflows without hard-coded workflow names."""
    file = Path(path)
    if not file.exists():
        raise FileNotFoundError(f"Workbook not found: {file}. Add the supplied Excel file or set WORKBOOK_PATH.")
    book = load_workbook(file, data_only=True)
    workflows: list[Workflow] = []
    for sheet in book.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [_field(c) for c in rows[0]]
        if any(headers):
            for row in rows[1:]:
                values = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
                item = _workflow(values, sheet.title)
                if item:
                    workflows.append(item)
        else:
            values = {_field(row[0]): row[1] for row in rows if len(row) >= 2 and _field(row[0])}
            item = _workflow(values, sheet.title)
            if item:
                workflows.append(item)
    unique = {workflow.id: workflow for workflow in workflows}
    if not unique:
        raise ValueError("No workflows found. Use recognised columns such as Workflow Name, Steps, Tools/APIs, Conditions and Expected Output.")
    return list(unique.values())
